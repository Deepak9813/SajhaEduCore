"""
Validator for multiple file types:
PDF, DOCX, XLSX, CSV, TXT
""" 

import csv
from pathlib import Path

from django.core.exceptions import ValidationError

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from apps.common.constants import (
    ALLOWED_FILE_EXTENSIONS,
    MAX_FILE_SIZE,
    ALLOWED_FILE_CONTENT_TYPES
    )


def validate_file(file):
    """Validate uploaded files."""

    if not file:
        # return file
        return

    # Validate file extension
    extension = Path(file.name).suffix.lower().lstrip(".")

    if extension not in ALLOWED_FILE_EXTENSIONS:
        raise ValidationError(
            "Unsupported file format."
        )

    # Validate file size
    if file.size > MAX_FILE_SIZE:
        raise ValidationError(
            "File size must not exceed 10 MB."
        )

    # Validate MIME type
    content_type = getattr(file, "content_type", None)

    if (
        content_type
        and content_type not in ALLOWED_FILE_CONTENT_TYPES
    ):
        raise ValidationError(
            "Invalid file type."
        )

    # Validate actual file content.
    # Prevents files with fake extensions (e.g., virus.exe renamed as .pdf, .docx, .xlsx, or .csv).
    try:
        if extension == "pdf":
            PdfReader(file)

        elif extension == "docx":
            Document(file)

        elif extension == "xlsx":
            load_workbook(file, read_only=True,)

        elif extension == "csv":
            decoded = file.read().decode("utf-8").splitlines()
            list(csv.reader(decoded))
            
        file.seek(0)  # Reset file pointer after reading

    except Exception:
        raise ValidationError(
            "Invalid file content."
        )
   